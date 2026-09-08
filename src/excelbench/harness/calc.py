"""Formula recalculation engines and oracle-backed comparison for ExcelBench."""

from __future__ import annotations

import json
import math
import shutil
import subprocess
import tempfile
from collections.abc import Mapping
from pathlib import Path
from typing import Any, Protocol

from openpyxl import load_workbook

from excelbench.harness.external_oracles import (
    ExternalOracleRequest,
    external_oracle_catalog,
    run_external_oracle,
)
from excelbench.results.calc_renderer import render_calc_results

SOFFICE_PATH = Path(
    shutil.which("soffice") or "/opt/homebrew/bin/soffice"
)


class CalcEngine(Protocol):
    """Protocol implemented by a spreadsheet formula recalculation engine."""

    name: str

    def available(self) -> bool:
        """Return whether this engine can run on the current machine."""

    def calculate(self, input_path: Path, output_path: Path) -> dict[str, Any]:
        """Calculate a workbook and return its formula-cell values."""


def formula_cells(input_path: Path) -> list[str]:
    """Return workbook-qualified references for every formula cell in a workbook."""
    workbook = load_workbook(input_path, data_only=False, read_only=True)
    try:
        return [
            f"{worksheet.title}!{cell.coordinate}"
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
    finally:
        workbook.close()


def libreoffice_version() -> str | None:
    """Return LibreOffice's version string when the configured binary is present."""
    if not SOFFICE_PATH.is_file():
        return None
    completed = subprocess.run(
        [str(SOFFICE_PATH), "--version"],
        capture_output=True,
        check=False,
        text=True,
        timeout=30,
    )
    if completed.returncode != 0:
        return None
    version = completed.stdout.strip() or completed.stderr.strip()
    return version or None


def recalculate_with_libreoffice(input_path: Path, output_path: Path) -> str | None:
    """Recalculate ``input_path`` through an isolated headless LibreOffice profile.

    The workbook is copied to a temporary source directory before ``--convert-to
    xlsx`` runs, so LibreOffice always writes a distinct file and its calculated
    cell caches can be copied to ``output_path`` safely.
    """
    if not SOFFICE_PATH.is_file():
        return f"LibreOffice binary not found at {SOFFICE_PATH}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="excelbench-libreoffice-") as temporary_dir:
        work_dir = Path(temporary_dir)
        source_dir = work_dir / "source"
        converted_dir = work_dir / "converted"
        profile_dir = work_dir / "profile"
        source_dir.mkdir()
        converted_dir.mkdir()
        profile_dir.mkdir()
        source_path = source_dir / input_path.name
        shutil.copy2(input_path, source_path)
        command = [
            str(SOFFICE_PATH),
            "--headless",
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(converted_dir),
            str(source_path),
        ]
        try:
            completed = subprocess.run(
                command,
                cwd=work_dir,
                capture_output=True,
                check=False,
                text=True,
                timeout=120,
            )
        except (OSError, subprocess.TimeoutExpired) as exc:
            return f"LibreOffice conversion could not run: {exc}"
        if completed.returncode != 0:
            detail = (completed.stderr or completed.stdout).strip() or "no diagnostic"
            return f"LibreOffice conversion failed: {detail}"

        converted_path = converted_dir / input_path.name
        if not converted_path.is_file():
            detail = (
                completed.stderr or completed.stdout
            ).strip() or "no output workbook"
            return f"LibreOffice conversion produced no XLSX: {detail}"
        shutil.copy2(converted_path, output_path)
    return None


def _read_values(
    workbook_path: Path, references: list[str], *, reader: str
) -> dict[str, Any]:
    if reader == "wolfxl":
        import wolfxl

        workbook = wolfxl.load_workbook(workbook_path, data_only=True)
        return {
            reference: workbook[sheet_name][cell_reference].value
            for reference in references
            for sheet_name, cell_reference in [reference.split("!", maxsplit=1)]
        }

    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        return {
            reference: workbook[sheet_name][cell_reference].value
            for reference in references
            for sheet_name, cell_reference in [reference.split("!", maxsplit=1)]
        }
    finally:
        workbook.close()


class WolfXLCalcEngine:
    """Formula calculation through WolfXL's public workbook API."""

    name = "wolfxl"

    def available(self) -> bool:
        """Return whether WolfXL can be imported."""
        try:
            import wolfxl  # noqa: F401
        except ImportError:
            return False
        return True

    def calculate(self, input_path: Path, output_path: Path) -> dict[str, Any]:
        """Calculate with WolfXL, save the workbook, and read cached results."""
        try:
            import wolfxl

            output_path.parent.mkdir(parents=True, exist_ok=True)
            references = formula_cells(input_path)
            workbook = wolfxl.load_workbook(input_path)
            workbook.calculate()
            workbook.save(output_path)
            return {
                "status": "passed",
                "values": _read_values(output_path, references, reader="wolfxl"),
                "reason": None,
            }
        except Exception as exc:  # Engine errors are data for the comparison report.
            return {
                "status": "failed",
                "values": {},
                "reason": f"{type(exc).__name__}: {exc}",
            }


class LibreOfficeCalcEngine:
    """Formula calculation through headless LibreOffice conversion."""

    name = "libreoffice"

    def available(self) -> bool:
        """Return whether the configured LibreOffice binary is available."""
        return libreoffice_version() is not None

    def calculate(self, input_path: Path, output_path: Path) -> dict[str, Any]:
        """Calculate with LibreOffice and read cached results from its output."""
        reason = recalculate_with_libreoffice(input_path, output_path)
        if reason is not None:
            return {"status": "failed", "values": {}, "reason": reason}
        try:
            return {
                "status": "passed",
                "values": _read_values(
                    output_path, formula_cells(input_path), reader="openpyxl"
                ),
                "reason": None,
            }
        except Exception as exc:  # Engine errors are data for the comparison report.
            return {
                "status": "failed",
                "values": {},
                "reason": f"{type(exc).__name__}: {exc}",
            }


class AsposeCellsFossCalcEngine:
    """Best-effort evaluation through the intentionally limited FOSS evaluator."""

    name = "aspose_cells_foss"

    def available(self) -> bool:
        """Return whether Aspose.Cells FOSS can be imported."""
        try:
            import aspose.cells_foss  # noqa: F401
            from aspose.cells_foss.formula_evaluator import FormulaEvaluator  # noqa: F401
        except ImportError:
            return False
        return True

    def calculate(self, input_path: Path, output_path: Path) -> dict[str, Any]:
        """Evaluate every formula with the FOSS evaluator and save its output."""
        try:
            from aspose.cells_foss import Workbook
            from aspose.cells_foss.formula_evaluator import FormulaEvaluator

            workbook = Workbook(str(input_path))
            evaluator = FormulaEvaluator(workbook)
            values: dict[str, Any] = {}
            unsupported: list[str] = []
            for worksheet in workbook.worksheets:
                for cell_reference, cell in worksheet.cells:
                    if not isinstance(cell.formula, str) or not cell.formula.startswith(
                        "="
                    ):
                        continue
                    value = evaluator.evaluate(cell.formula, worksheet, cell_reference)
                    reference = f"{worksheet.name}!{cell_reference}"
                    values[reference] = value
                    if value is None:
                        unsupported.append(reference)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            workbook.save(str(output_path))
            if unsupported:
                preview = ", ".join(unsupported[:5])
                return {
                    "status": "failed",
                    "values": values,
                    "reason": (
                        "FOSS FormulaEvaluator returned no value for "
                        f"{len(unsupported)} formulas: {preview}"
                    ),
                }
            return {"status": "passed", "values": values, "reason": None}
        except Exception as exc:  # Its unsupported parser must not escape the suite.
            return {
                "status": "failed",
                "values": {},
                "reason": f"{type(exc).__name__}: {exc}",
            }


class ZavoraCalcEngine:
    """Formula calculation through the optional Zavora external oracle helper."""

    name = "zavora"

    def available(self) -> bool:
        """Return whether the catalogued Zavora helper can be launched."""
        return external_oracle_catalog(_repository_root())["zavora"].is_available()

    def calculate(self, input_path: Path, output_path: Path) -> dict[str, Any]:
        """Delegate calculation to Zavora's JSON external-oracle protocol."""
        tool = external_oracle_catalog(_repository_root())["zavora"]
        if not tool.is_available():
            return {
                "status": "unavailable",
                "values": {},
                "reason": "Zavora helper is unavailable",
            }
        result = run_external_oracle(
            tool,
            ExternalOracleRequest(
                fixture_id=input_path.stem,
                operation="calculate",
                payload={},
                input_path=input_path.resolve(),
                output_path=output_path.resolve(),
            ),
        )
        if result.skipped:
            return {
                "status": "unavailable",
                "values": {},
                "reason": result.notes or "Zavora helper skipped",
            }
        if not result.passed:
            detail = (
                result.notes
                or result.stderr.strip()
                or result.stdout.strip()
                or "no diagnostic"
            )
            return {
                "status": "failed",
                "values": {},
                "reason": f"Zavora calculation failed: {detail}",
            }
        values = result.payload.get("values")
        if not isinstance(values, dict):
            return {
                "status": "failed",
                "values": {},
                "reason": "Zavora response omitted values",
            }
        return {"status": "passed", "values": values, "reason": None}


def values_match(expected: Any, actual: Any) -> bool:
    """Compare calculation values with number tolerance and exact scalar semantics."""
    if isinstance(expected, bool) or isinstance(actual, bool):
        return (
            isinstance(expected, bool)
            and isinstance(actual, bool)
            and expected is actual
        )
    if isinstance(expected, (int, float)) and isinstance(actual, (int, float)):
        return math.isclose(float(expected), float(actual), abs_tol=1e-6, rel_tol=1e-9)
    return bool(expected == actual)


def run_calc_suite(
    fixture: Path, expected: dict[str, Any] | Path, output_dir: Path
) -> dict[str, Any]:
    """Run available engines against the LibreOffice calculation oracle fixture."""
    expected_data = _load_expected(expected)
    expected_cells = expected_data["cells"]
    if not isinstance(expected_cells, Mapping):
        raise ValueError("Expected calculation data must contain a cells mapping")

    output_dir.mkdir(parents=True, exist_ok=True)
    engine_results: dict[str, dict[str, Any]] = {}
    for engine in _engines():
        if not engine.available():
            engine_results[engine.name] = {
                "status": "unavailable",
                "matched": 0,
                "total": len(expected_cells),
                "mismatched_cells": [],
                "reason": f"{engine.name} is unavailable",
            }
            continue
        engine_result = engine.calculate(
            fixture, output_dir / engine.name / fixture.name
        )
        values = engine_result.get("values", {})
        if not isinstance(values, Mapping):
            values = {}
        matched, mismatches = _compare_expected_values(expected_cells, values)
        total = len(expected_cells)
        reason: str | None
        if engine_result["status"] == "passed" and matched < total:
            # An engine that runs but produces wrong or missing values is a
            # failed calculation run, not a pass.
            status = "failed"
            missing = sum(1 for m in mismatches if m.get("actual") is None)
            wrong = len(mismatches) - missing
            reason = (
                f"{total - matched}/{total} formula values wrong or missing "
                f"(first mismatch batch: {len(mismatches)} shown, "
                f"{missing} missing, {wrong} wrong)"
            )
        else:
            status = str(engine_result["status"])
            reason = engine_result.get("reason")
        engine_results[engine.name] = {
            "status": status,
            "matched": matched,
            "total": total,
            "mismatched_cells": mismatches,
            "reason": reason,
        }

    results = {
        "fixture": str(fixture),
        "oracle": expected_data.get("oracle"),
        "engines": engine_results,
    }
    render_calc_results(results, output_dir)
    return results


def _compare_expected_values(
    expected_cells: Mapping[str, Any], actual_values: Mapping[str, Any]
) -> tuple[int, list[dict[str, Any]]]:
    matched = 0
    mismatches: list[dict[str, Any]] = []
    for reference, expected_cell in expected_cells.items():
        if not isinstance(expected_cell, Mapping) or "value" not in expected_cell:
            raise ValueError(f"Expected cell {reference} has no value")
        expected_value = expected_cell["value"]
        actual_value = actual_values.get(reference)
        if values_match(expected_value, actual_value):
            matched += 1
        elif len(mismatches) < 10:
            mismatches.append(
                {"cell": reference, "expected": expected_value, "actual": actual_value}
            )
    return matched, mismatches


def _engines() -> tuple[CalcEngine, ...]:
    return (
        WolfXLCalcEngine(),
        LibreOfficeCalcEngine(),
        AsposeCellsFossCalcEngine(),
        ZavoraCalcEngine(),
    )


def _load_expected(expected: dict[str, Any] | Path) -> dict[str, Any]:
    if isinstance(expected, Path):
        expected_path = (
            expected / "expected_values.json" if expected.is_dir() else expected
        )
        loaded: dict[str, Any] = json.loads(expected_path.read_text())
        return loaded
    return expected


def _repository_root() -> Path:
    return Path(__file__).resolve().parents[3]
