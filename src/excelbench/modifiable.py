"""Workbook engines used by the surgical mutation benchmark."""

from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any, Protocol

from excelbench.harness.external_oracles import (
    ExternalOracleRequest,
    external_oracle_catalog,
    run_external_oracle,
)

try:
    import openpyxl as _openpyxl
except ImportError:  # pragma: no cover - optional dependency guard
    _openpyxl = None

try:
    import wolfxl as _wolfxl
except ImportError:  # pragma: no cover - optional dependency guard
    _wolfxl = None

try:
    import aspose.cells_foss as _aspose_cells_foss
except ImportError:  # pragma: no cover - optional dependency guard
    _aspose_cells_foss = None

Mutation = dict[str, Any]
_REPO_ROOT = Path(__file__).resolve().parents[2]


class ModifiableEngine(Protocol):
    """Protocol implemented by engines that can modify an existing workbook."""

    name: str

    def available(self) -> bool:
        """Return whether this engine is installed and usable."""

    def mutate(self, template: Path, output: Path, mutations: list[Mutation]) -> None:
        """Apply all mutations to a template workbook and save the output."""


def _mutation_target(mutation: Mutation) -> tuple[str, str, Any]:
    """Validate and unpack one mutation payload."""
    try:
        sheet = mutation["sheet"]
        cell = mutation["cell"]
        value = mutation["value"]
    except KeyError as error:
        raise ValueError(
            f"Mutation is missing required key: {error.args[0]}"
        ) from error
    if not isinstance(sheet, str) or not sheet:
        raise ValueError("Mutation sheet must be a non-empty string")
    if not isinstance(cell, str) or not cell:
        raise ValueError("Mutation cell must be a non-empty string")
    return sheet, cell, value


class OpenpyxlEngine:
    """In-place workbook mutation through OpenPyXL."""

    name = "openpyxl"

    def available(self) -> bool:
        """Return whether OpenPyXL can be imported."""
        return _openpyxl is not None

    def mutate(self, template: Path, output: Path, mutations: list[Mutation]) -> None:
        """Apply mutations with OpenPyXL and save a new workbook package."""
        if _openpyxl is None:
            raise RuntimeError("openpyxl is not installed")
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook = _openpyxl.load_workbook(template)
        try:
            for mutation in mutations:
                sheet, cell, value = _mutation_target(mutation)
                workbook[sheet][cell] = value
            workbook.save(output)
        finally:
            workbook.close()


class WolfxlEngine:
    """In-place workbook mutation through WolfXL modify mode."""

    name = "wolfxl"

    def available(self) -> bool:
        """Return whether WolfXL can be imported."""
        return _wolfxl is not None

    def mutate(self, template: Path, output: Path, mutations: list[Mutation]) -> None:
        """Apply mutations with WolfXL's public modify-mode API."""
        if _wolfxl is None:
            raise RuntimeError("wolfxl is not installed")
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook = _wolfxl.load_workbook(template, modify=True)
        for mutation in mutations:
            sheet, cell, value = _mutation_target(mutation)
            workbook[sheet][cell].value = value
        workbook.save(output)


class AsposeFossEngine:
    """In-place workbook mutation through aspose-cells-foss."""

    name = "aspose-cells-foss"

    def available(self) -> bool:
        """Return whether aspose-cells-foss can be imported."""
        return _aspose_cells_foss is not None

    def mutate(self, template: Path, output: Path, mutations: list[Mutation]) -> None:
        """Apply mutations with aspose-cells-foss and save a new workbook package."""
        if _aspose_cells_foss is None:
            raise RuntimeError("aspose-cells-foss is not installed")
        output.parent.mkdir(parents=True, exist_ok=True)
        workbook = _aspose_cells_foss.Workbook(str(template))
        for mutation in mutations:
            sheet, cell, value = _mutation_target(mutation)
            worksheet = next(
                (
                    candidate
                    for candidate in workbook.worksheets
                    if candidate.name == sheet
                ),
                None,
            )
            if worksheet is None:
                raise KeyError(f"Worksheet not found: {sheet}")
            worksheet.cells.get_cell_by_name(cell).value = value
        workbook.save(str(output))


class ZavoraEngine:
    """In-place workbook mutation through the optional Zavora oracle helper."""

    name = "zavora-xlsx"

    def available(self) -> bool:
        """Return whether the Zavora helper command and source are available."""
        return external_oracle_catalog(_REPO_ROOT)["zavora"].is_available()

    def mutate(self, template: Path, output: Path, mutations: list[Mutation]) -> None:
        """Apply mutations through chained external-oracle requests."""
        tool = external_oracle_catalog(_REPO_ROOT)["zavora"]
        if not tool.is_available():
            raise RuntimeError("zavora external oracle is unavailable")
        output.parent.mkdir(parents=True, exist_ok=True)
        if not mutations:
            shutil.copyfile(template, output)
            return

        intermediate_paths: list[Path] = []
        input_path = template.resolve()
        try:
            for index, mutation in enumerate(mutations):
                sheet, cell, value = _mutation_target(mutation)
                is_last = index == len(mutations) - 1
                next_path = output.resolve()
                if not is_last:
                    next_path = output.with_name(
                        f".{output.stem}.mutation-{index}{output.suffix}"
                    ).resolve()
                    intermediate_paths.append(next_path)
                request = ExternalOracleRequest(
                    fixture_id="mutation-template",
                    operation="mutate",
                    payload={"sheet": sheet, "cell": cell, "value": value},
                    input_path=input_path,
                    output_path=next_path,
                )
                result = run_external_oracle(tool, request)
                if not result.passed:
                    detail = (
                        result.notes
                        or result.stderr
                        or result.stdout
                        or "unknown oracle failure"
                    )
                    raise RuntimeError(f"zavora mutation failed: {detail.strip()}")
                input_path = next_path
        finally:
            for intermediate_path in intermediate_paths:
                intermediate_path.unlink(missing_ok=True)


def modifiable_engines() -> list[ModifiableEngine]:
    """Return the fixed engine registry used by the mutation suite."""
    return [OpenpyxlEngine(), WolfxlEngine(), AsposeFossEngine(), ZavoraEngine()]
