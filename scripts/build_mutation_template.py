#!/usr/bin/env python3
"""Build the deterministic mutation-fidelity workbook fixture."""

from __future__ import annotations

import hashlib
import json
import tempfile
import xml.etree.ElementTree as ET
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE_DIR = REPO_ROOT / "fixtures" / "mutation"
FIXTURE_PATH = FIXTURE_DIR / "template_corporate_model.xlsx"
MANIFEST_PATH = FIXTURE_DIR / "manifest.json"
FIXED_TIMESTAMP = (2026, 1, 1, 0, 0, 0)
GENERATED_AT = "2026-01-01T00:00:00Z"
_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_OFFICE_DOCUMENT_REL_NS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
)


def _anchor(from_cell: str, to_cell: str) -> TwoCellAnchor:
    """Create the explicit two-cell anchor used by the OpenPyXL adapter."""
    from_col = ord(from_cell[0].upper()) - ord("A")
    from_row = int(from_cell[1:]) - 1
    to_col = ord(to_cell[0].upper()) - ord("A")
    to_row = int(to_cell[1:]) - 1
    return TwoCellAnchor(
        _from=AnchorMarker(col=from_col, colOff=0, row=from_row, rowOff=0),
        to=AnchorMarker(col=to_col, colOff=0, row=to_row, rowOff=0),
    )


def _make_workbook(path: Path) -> None:
    """Build the structural workbook before custom XML injection."""
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    schedule = workbook.create_sheet("Schedule")
    summary = workbook.create_sheet("Summary")
    workbook.properties.created = "2026-01-01T00:00:00Z"
    workbook.properties.modified = "2026-01-01T00:00:00Z"

    headers = ["Driver", "Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7"]
    inputs.append(headers)
    for row_index in range(2, 42):
        inputs.cell(row_index, 1, f"Input {row_index - 1:02d}")
        for column_index in range(2, 9):
            inputs.cell(row_index, column_index, (row_index - 1) * column_index * 10)
    for cell in inputs[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for column_index in range(1, 9):
        inputs.column_dimensions[get_column_letter(column_index)].width = 14
    inputs.freeze_panes = "B2"
    inputs["A2"].hyperlink = "https://example.com/corporate-model-inputs"
    inputs["A2"].comment = Comment("Mutation target workbook input.", "ExcelBench")

    inputs_table = Table(displayName="CorporateInputs", ref="A1:H41")
    inputs_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    inputs.add_table(inputs_table)

    list_validation = DataValidation(
        type="list", formula1='"Low,Medium,High"', allow_blank=True
    )
    list_validation.add("H2:H41")
    inputs.add_data_validation(list_validation)
    decimal_validation = DataValidation(
        type="decimal",
        operator="between",
        formula1="0",
        formula2="100000",
        allow_blank=False,
    )
    decimal_validation.add("B2:G41")
    inputs.add_data_validation(decimal_validation)
    inputs.conditional_formatting.add(
        "B2:H41",
        CellIsRule(
            operator="greaterThan",
            formula=["1000"],
            fill=PatternFill("solid", fgColor="FFF2CC"),
        ),
    )
    inputs.conditional_formatting.add(
        "B2:H41",
        ColorScaleRule(
            start_type="min",
            start_color="F8696B",
            mid_type="percentile",
            mid_value=50,
            mid_color="FFEB84",
            end_type="max",
            end_color="63BE7B",
        ),
    )

    schedule.append(["Schedule Metric", *headers[1:]])
    for row_index in range(2, 15):
        schedule.cell(row_index, 1, f"Schedule {row_index - 1:02d}")
        for column_index in range(2, 9):
            column_letter = get_column_letter(column_index)
            input_row = row_index + 1
            if row_index == 2:
                formula = f"=Inputs!{column_letter}{input_row}"
            elif row_index == 3:
                formula = f"=SUM({column_letter}2,Inputs!{column_letter}{input_row})"
            else:
                formula = f"=SUM({column_letter}2:{column_letter}{row_index - 1})"
            schedule.cell(row_index, column_index, formula)
    schedule.freeze_panes = "B2"
    schedule.protection.sheet = True

    summary.append(["Summary Metric", *headers[1:]])
    for row_index in range(2, 12):
        summary.cell(row_index, 1, f"Summary {row_index - 1:02d}")
        source_row = row_index + 2
        for column_index in range(2, 9):
            column_letter = get_column_letter(column_index)
            summary.cell(
                row_index, column_index, f"=Schedule!{column_letter}{source_row}"
            )
    summary.freeze_panes = "B2"
    summary.page_setup.orientation = "landscape"
    summary.page_setup.fitToWidth = 1
    summary.print_title_rows = "$1:$2"
    summary.oddHeader.center.text = "Quarterly Model"

    values = Reference(summary, min_col=2, max_col=8, min_row=1, max_row=11)
    categories = Reference(summary, min_col=1, min_row=2, max_row=11)
    bar_chart = BarChart()
    bar_chart.title = "Quarterly Summary"
    bar_chart.add_data(values, titles_from_data=True)
    bar_chart.set_categories(categories)
    bar_chart.anchor = _anchor("B2", "H12")
    summary.add_chart(bar_chart)

    line_chart = LineChart()
    line_chart.title = "Quarterly Trend"
    line_chart.add_data(values, titles_from_data=True)
    line_chart.set_categories(categories)
    line_chart.anchor = _anchor("J2", "P12")
    summary.add_chart(line_chart)

    workbook.defined_names.add(
        DefinedName("CorporateInputs", attr_text="'Inputs'!$B$2:$H$10")
    )
    workbook.defined_names.add(
        DefinedName("LocalGrowthRate", attr_text="'Inputs'!$B$2", localSheetId=0)
    )
    workbook.save(path)
    workbook.close()


def _inject_custom_xml(path: Path) -> None:
    """Inject package parts OpenPyXL intentionally does not preserve on save."""
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}

    root_rels = ET.fromstring(parts["_rels/.rels"])
    ET.register_namespace("", _RELATIONSHIPS_NS)
    relation_tag = f"{{{_RELATIONSHIPS_NS}}}Relationship"
    if not any(
        relation.get("Type") == f"{_OFFICE_DOCUMENT_REL_NS}/customXml"
        for relation in root_rels
    ):
        existing_ids = {relation.get("Id") for relation in root_rels}
        relation_id = "rId1"
        suffix = 1
        while relation_id in existing_ids:
            suffix += 1
            relation_id = f"rId{suffix}"
        ET.SubElement(
            root_rels,
            relation_tag,
            {
                "Id": relation_id,
                "Type": f"{_OFFICE_DOCUMENT_REL_NS}/customXml",
                "Target": "customXml/item1.xml",
            },
        )
    parts["_rels/.rels"] = ET.tostring(
        root_rels, encoding="utf-8", xml_declaration=True
    )

    content_types = ET.fromstring(parts["[Content_Types].xml"])
    ET.register_namespace("", _CONTENT_TYPES_NS)
    override_tag = f"{{{_CONTENT_TYPES_NS}}}Override"
    required_override = "/customXml/itemProps1.xml"
    if not any(item.get("PartName") == required_override for item in content_types):
        ET.SubElement(
            content_types,
            override_tag,
            {
                "PartName": required_override,
                "ContentType": (
                    "application/vnd.openxmlformats-officedocument.customXmlProperties+xml"
                ),
            },
        )
    parts["[Content_Types].xml"] = ET.tostring(
        content_types, encoding="utf-8", xml_declaration=True
    )

    parts["customXml/item1.xml"] = b"""<?xml version=\"1.0\" encoding=\"UTF-8\"?>
<wolfxl:p1Fidelity xmlns:wolfxl=\"https://wolfxl.local/fidelity\">
  <wolfxl:source>openpyxl-targeted-fixture</wolfxl:source>
</wolfxl:p1Fidelity>"""
    parts["customXml/itemProps1.xml"] = b"""<?xml version=\"1.0\" encoding=\"UTF-8\"?>
<ds:datastoreItem ds:itemID=\"{11111111-2222-3333-4444-555555555555}\"
  xmlns:ds=\"http://schemas.openxmlformats.org/officeDocument/2006/customXml\">
  <ds:schemaRefs/>
</ds:datastoreItem>"""
    parts[
        "customXml/_rels/item1.xml.rels"
    ] = b"""<?xml version=\"1.0\" encoding=\"UTF-8\"?>
<Relationships xmlns=\"http://schemas.openxmlformats.org/package/2006/relationships\">
  <Relationship Id=\"rId1\"
    Type=\"http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps\"
    Target=\"itemProps1.xml\"/>
</Relationships>"""

    with ZipFile(path, "w", compression=ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(parts):
            info = ZipInfo(filename=name, date_time=FIXED_TIMESTAMP)
            info.compress_type = ZIP_DEFLATED
            archive.writestr(info, parts[name])


def _sha256(data: bytes) -> str:
    """Return a SHA-256 digest for fixture package content."""
    return hashlib.sha256(data).hexdigest()


def build_fixture() -> Path:
    """Build the workbook and manifest, returning the workbook path."""
    FIXTURE_DIR.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(
        suffix=".xlsx", dir=FIXTURE_DIR, delete=False
    ) as temporary:
        staging_path = Path(temporary.name)
    try:
        _make_workbook(staging_path)
        _inject_custom_xml(staging_path)
        staging_path.replace(FIXTURE_PATH)
    finally:
        staging_path.unlink(missing_ok=True)

    with ZipFile(FIXTURE_PATH) as archive:
        part_names = sorted(archive.namelist())
        checksums = {name: _sha256(archive.read(name)) for name in part_names}
    manifest = {
        "generated_at": GENERATED_AT,
        "generator": "openpyxl-structural+customxml-injection",
        "parts": part_names,
        "checksums": checksums,
        "mutation": {"sheet": "Inputs", "cell": "B2", "value": "MUTATED-<engine>"},
        "second_mutation": {"sheet": "Inputs", "cell": "C3", "value": 1337},
    }
    MANIFEST_PATH.write_text(
        json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return FIXTURE_PATH


def main() -> None:
    """Build the committed mutation fixture."""
    print(build_fixture())


if __name__ == "__main__":
    main()
