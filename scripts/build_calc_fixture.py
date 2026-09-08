#!/usr/bin/env python3
"""Build the deterministic Formula Recalculation benchmark fixture.

The canonical fixture is written by openpyxl and ships formulas with NO cached
values. LibreOffice's isolated, headless ``--convert-to xlsx`` path recalculates
a throwaway copy; every formula cache in that copy is verified before the
oracle JSON is written. Engines are scored against the oracle, so a passing
engine must produce the values itself rather than pass stored caches through.
"""

from __future__ import annotations

import json
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils.datetime import to_excel

from excelbench.harness.calc import libreoffice_version, recalculate_with_libreoffice

REPOSITORY_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_FIXTURE_PATH = REPOSITORY_ROOT / "fixtures" / "calc" / "financial_model.xlsx"
DEFAULT_EXPECTED_PATH = REPOSITORY_ROOT / "fixtures" / "calc" / "expected_values.json"


def build_fixture(
    fixture_path: Path = DEFAULT_FIXTURE_PATH,
    expected_path: Path = DEFAULT_EXPECTED_PATH,
) -> dict[str, Any]:
    """Build a cache-free workbook, derive oracle values via LibreOffice, persist both.

    The canonical fixture keeps NO cached formula values: openpyxl writes the
    formulas only. Expected values are harvested from a throwaway LibreOffice
    recalculation copy. A passing engine must therefore produce the numbers
    itself; passthrough of precomputed caches cannot score.
    """
    fixture_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = _build_workbook()
    workbook.save(fixture_path)

    formula_references = _formula_references(fixture_path)
    with tempfile.TemporaryDirectory(prefix="excelbench-calc-oracle-") as temporary_dir:
        recalculated_path = Path(temporary_dir) / "recalculated.xlsx"
        recalculate_reason = recalculate_with_libreoffice(
            fixture_path, recalculated_path
        )
        if recalculate_reason is not None:
            raise RuntimeError(recalculate_reason)
        cells = _cached_cells(recalculated_path, formula_references)
    _assert_fixture_cache_free(fixture_path, formula_references)

    expected = {
        "oracle": f"libreoffice/{libreoffice_version() or 'unknown'}",
        "cells": cells,
    }
    expected_path.parent.mkdir(parents=True, exist_ok=True)
    expected_path.write_text(json.dumps(expected, indent=2, sort_keys=True) + "\n")
    return expected


def _assert_fixture_cache_free(
    fixture_path: Path, formula_references: list[str]
) -> None:
    """Fail loudly if the canonical fixture ever ships cached formula values."""
    workbook = load_workbook(fixture_path, data_only=True, read_only=True)
    try:
        cached = [
            reference
            for reference in formula_references
            if workbook[reference.split("!", maxsplit=1)[0]][
                reference.split("!", maxsplit=1)[1]
            ].value
            is not None
        ]
    finally:
        workbook.close()
    if cached:
        sample = ", ".join(sorted(cached)[:5])
        raise RuntimeError(
            f"fixture must not ship cached formula values (found: {sample})"
        )


def _build_workbook() -> Workbook:
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    schedule = workbook.create_sheet("Schedule")
    summary = workbook.create_sheet("Summary")
    _populate_inputs(inputs)
    _populate_schedule(schedule)
    _populate_summary(summary)
    return workbook


def _populate_inputs(worksheet: Any) -> None:
    worksheet["A1"] = "Financial model assumptions"
    assumptions: list[tuple[str, Any]] = [
        ("Start Date", date(2025, 1, 31)),
        ("Revenue Growth", 0.045),
        ("Gross Margin", 0.62),
        ("Fixed Costs", 12000.0),
        ("Tax Rate", 0.24),
        ("Discount Rate", 0.1),
        ("Target Category", "Enterprise"),
        ("Target Region", "West"),
    ]
    for row_index, (label, value) in enumerate(assumptions, start=2):
        worksheet.cell(row=row_index, column=1, value=label)
        worksheet.cell(row=row_index, column=2, value=value)
    worksheet["B2"].number_format = "yyyy-mm-dd"
    worksheet["B3"].number_format = "0.0%"
    worksheet["B4"].number_format = "0.0%"
    worksheet["B6"].number_format = "0.0%"
    worksheet["B7"].number_format = "0.0%"

    headers = ["Record", "Category", "Units", "Price", "Launch Date", "Region"]
    for column_index, header in enumerate(headers, start=1):
        worksheet.cell(row=12, column=column_index, value=header)
    categories = ("Enterprise", "SMB", "Consumer", "Enterprise", "SMB")
    regions = ("West", "East", "South", "North")
    for offset in range(40):
        row_index = offset + 13
        worksheet.cell(row=row_index, column=1, value=offset + 1)
        worksheet.cell(
            row=row_index, column=2, value=categories[offset % len(categories)]
        )
        worksheet.cell(row=row_index, column=3, value=20 + ((offset * 7) % 35))
        worksheet.cell(
            row=row_index, column=4, value=125.0 + ((offset * 13) % 170) + 0.25
        )
        worksheet.cell(row=row_index, column=5, value=date(2025, (offset % 12) + 1, 1))
        worksheet.cell(row=row_index, column=6, value=regions[offset % len(regions)])


def _populate_schedule(worksheet: Any) -> None:
    worksheet["A1"] = "Metric"
    for period in range(1, 13):
        worksheet.cell(row=1, column=period + 1, value=period)
    labels = [
        "Revenue",
        "Cost of Goods Sold",
        "Gross Profit",
        "Fixed Costs",
        "Target Units",
        "Target Regional Price",
        "Operating Costs",
        "EBITDA",
        "Tax",
        "Cash Flow",
    ]
    for row_index, label in enumerate(labels, start=2):
        worksheet.cell(row=row_index, column=1, value=label)
    for column_index in range(2, 14):
        column_letter = worksheet.cell(row=1, column=column_index).column_letter
        formulas = {
            2: f"=100000*(1+Inputs!$B$3)^({column_letter}$1-1)",
            3: f"={column_letter}2*(1-Inputs!$B$4)",
            4: f"={column_letter}2-{column_letter}3",
            5: "=Inputs!$B$5",
            6: "=SUMIF(Inputs!$B$13:$B$52,Inputs!$B$8,Inputs!$C$13:$C$52)",
            7: "=SUMIFS(Inputs!$D$13:$D$52,Inputs!$B$13:$B$52,Inputs!$B$8,"
            "Inputs!$F$13:$F$52,Inputs!$B$9)",
            8: f"={column_letter}3+{column_letter}5+{column_letter}6*{column_letter}7+SUM(0)",
            9: f"={column_letter}2-{column_letter}8",
            10: f"=IF({column_letter}9>0,{column_letter}9*Inputs!$B$6,0)",
            11: f"={column_letter}9-{column_letter}10-{column_letter}2*0.08",
        }
        for row_index, formula in formulas.items():
            worksheet.cell(row=row_index, column=column_index, value=formula)


def _populate_summary(worksheet: Any) -> None:
    worksheet["A1"] = "Summary Metric"
    worksheet["B1"] = "Value"
    formulas = {
        "A2": "Positive cash flow flag",
        "B2": "=IF(SUM(Schedule!B11:M11)>0,1,0)",
        "A3": "Growth tier",
        "B3": "=IF(Inputs!$B$3>=0.05,3,IF(Inputs!$B$3>=0.03,2,1))",
        "A4": "Minimum EBITDA",
        "B4": "=MIN(Schedule!B9:M9)",
        "A5": "Maximum EBITDA",
        "B5": "=MAX(Schedule!B9:M9)",
        "A6": "Average cash flow",
        "B6": "=ROUND(AVERAGE(Schedule!B11:M11),2)",
        "A7": "VLOOKUP enterprise units",
        "B7": '=VLOOKUP("Enterprise",Inputs!$B$13:$D$52,2,FALSE)',
        "A8": "INDEX/MATCH enterprise price",
        "B8": '=INDEX(Inputs!$D$13:$D$52,MATCH("Enterprise",Inputs!$B$13:$B$52,0))',
        "A10": "Net present value",
        "B10": "=NPV(Inputs!$B$7,Schedule!B11:M11)",
        "A11": "Internal rate of return",
        "B11": "=IRR(B18:N18)",
        "A12": "Date plus three months",
        "B12": "=EDATE(Inputs!$B$2,3)",
        "A13": "Month end plus three months",
        "B13": "=EOMONTH(Inputs!$B$2,3)",
        "A14": "Enterprise records",
        "B14": '=COUNTIF(Inputs!$B$13:$B$52,"Enterprise")',
        "A15": "Enterprise west records",
        "B15": '=COUNTIFS(Inputs!$B$13:$B$52,"Enterprise",Inputs!$F$13:$F$52,"West")',
    }
    for reference, value in formulas.items():
        worksheet[reference] = value

    worksheet["A17"] = "Signed cash-flow series for IRR"
    cash_flows = [
        -100000.0,
        16500.0,
        18000.0,
        20500.0,
        23000.0,
        25000.0,
        27000.0,
        29500.0,
        31500.0,
        34000.0,
        36000.0,
        38500.0,
        41000.0,
    ]
    for column_index, cash_flow in enumerate(cash_flows, start=2):
        worksheet.cell(row=18, column=column_index, value=cash_flow)


def _formula_references(fixture_path: Path) -> list[str]:
    workbook = load_workbook(fixture_path, data_only=False, read_only=True)
    try:
        return [
            f"{worksheet.title}!{cell.coordinate}"
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows()
            for cell in row
            if isinstance(cell.value, str) and cell.value.startswith("=")
        ]
    finally:
        workbook.close()


def _cached_cells(
    fixture_path: Path, formula_references: list[str]
) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(fixture_path, data_only=True, read_only=True)
    try:
        cells: dict[str, dict[str, Any]] = {}
        for reference in formula_references:
            sheet_name, cell_reference = reference.split("!", maxsplit=1)
            value = workbook[sheet_name][cell_reference].value
            if value is None:
                raise RuntimeError(
                    f"LibreOffice left an empty cached formula value at {reference}"
                )
            cells[reference] = _expected_cell(value)
        return cells
    finally:
        workbook.close()


def _expected_cell(value: Any) -> dict[str, Any]:
    if isinstance(value, bool):
        return {"value": value, "type": "bool"}
    if isinstance(value, (int, float)):
        return {"value": round(float(value), 10), "type": "number"}
    if isinstance(value, (date, datetime)):
        return {"value": round(float(to_excel(value)), 10), "type": "number"}
    if isinstance(value, str):
        return {"value": value, "type": "string"}
    raise TypeError(f"Unsupported cached formula value: {value!r}")


def main() -> int:
    """Build the fixture and report formula coverage with representative values."""
    expected = build_fixture()
    cells = expected["cells"]
    print(f"Formula cells: {len(cells)}")
    for reference in sorted(cells)[:3]:
        print(f"{reference}: {cells[reference]}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
