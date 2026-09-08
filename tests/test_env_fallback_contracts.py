"""Env-independent contracts for optional-dependency harness paths.

Covers calc-tier status semantics, mutation-suite control flow, modifiable
engine guards, renderers, and the Zavora adapter's payload/error paths on
machines without LibreOffice, aspose-cells-foss, or the Rust helper.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import Any

import openpyxl
import pytest

from excelbench import modifiable as modifiable_module
from excelbench.harness import calc as calc_module
from excelbench.harness import mutation as mutation_module
from excelbench.harness.adapters import zavora_adapter as zavora_module
from excelbench.harness.adapters.zavora_adapter import ZavoraAdapter
from excelbench.harness.calc import (
    AsposeCellsFossCalcEngine,
    LibreOfficeCalcEngine,
    ZavoraCalcEngine,
    formula_cells,
    run_calc_suite,
)
from excelbench.models import CellType, CellValue
from excelbench.modifiable import OpenpyxlEngine, _mutation_target
from excelbench.results.mutation_renderer import render_mutation_report

REPO_ROOT = Path(__file__).resolve().parents[1]
FIXTURE = REPO_ROOT / "fixtures" / "calc" / "financial_model.xlsx"


class _FakeEngine:
    def __init__(
        self,
        name: str,
        available: bool = True,
        status: str = "passed",
        values: dict[str, Any] | None = None,
        reason: str | None = None,
        raises: Exception | None = None,
    ) -> None:
        self.name = name
        self._available = available
        self._status = status
        self._values = values or {}
        self._reason = reason
        self._raises = raises

    def available(self) -> bool:
        return self._available

    def calculate(self, input_path: Path, output_path: Path) -> dict[str, Any]:
        if self._raises is not None:
            raise self._raises
        return {"status": self._status, "values": self._values, "reason": self._reason}


def _expected(cells: dict[str, Any]) -> dict[str, Any]:
    return {"oracle": "test-oracle", "cells": cells}


def _cell(value: Any, cell_type: str = "number") -> dict[str, Any]:
    return {"value": value, "type": cell_type}


# --------------------------------------------------------------------------
# Calc-tier status semantics
# --------------------------------------------------------------------------


def test_calc_suite_marks_partial_match_failed_with_missing_wrong_split(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    values = {
        "Sheet1!A1": 1,
        "Sheet1!A2": None,
        "Sheet1!A3": 99,
    }
    monkeypatch.setattr(
        calc_module, "_engines", lambda: (_FakeEngine("partial", values=values),)
    )
    results = run_calc_suite(
        FIXTURE,
        _expected(
            {
                "Sheet1!A1": _cell(1),
                "Sheet1!A2": _cell(2),
                "Sheet1!A3": _cell(3),
                "Sheet1!A4": _cell(4),
            }
        ),
        tmp_path,
    )
    engine = results["engines"]["partial"]
    assert engine["status"] == "failed"
    assert engine["matched"] == 1
    assert engine["total"] == 4
    assert "3/4 formula values wrong or missing" in engine["reason"]
    assert "2 missing" in engine["reason"]
    assert "1 wrong" in engine["reason"]


def test_calc_suite_full_match_passes_and_renders(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        calc_module,
        "_engines",
        lambda: (
            _FakeEngine("good", values={"Sheet1!A1": 10}),
            _FakeEngine("bad", status="failed", values={}, reason="exploded"),
            _FakeEngine(
                "absent",
                available=False,
                values={"Sheet1!A1": 10},
            ),
        ),
    )
    results = run_calc_suite(FIXTURE, _expected({"Sheet1!A1": _cell(10)}), tmp_path)
    assert results["oracle"] == "test-oracle"
    assert results["engines"]["good"]["status"] == "passed"
    assert results["engines"]["good"]["reason"] is None
    assert results["engines"]["bad"]["status"] == "failed"
    assert results["engines"]["bad"]["reason"] == "exploded"
    assert results["engines"]["absent"]["status"] == "unavailable"
    assert results["engines"]["absent"]["reason"] == "absent is unavailable"

    rendered = json.loads((tmp_path / "results.json").read_text())
    assert set(rendered["engines"]) == {"good", "bad", "absent"}
    readme = (tmp_path / "README.md").read_text()
    assert "good" in readme and "unavailable" in readme.lower()


def test_calc_suite_values_must_be_mapping(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    class _ListValuesEngine(_FakeEngine):
        def calculate(self, input_path: Path, output_path: Path) -> dict[str, Any]:
            result = super().calculate(input_path, output_path)
            result["values"] = ["not", "a", "mapping"]
            return result

    monkeypatch.setattr(calc_module, "_engines", lambda: (_ListValuesEngine("listy"),))
    results = run_calc_suite(FIXTURE, _expected({"A1": _cell(1)}), tmp_path)
    assert results["engines"]["listy"]["status"] == "failed"
    assert results["engines"]["listy"]["matched"] == 0


def test_calc_suite_rejects_malformed_expected_payloads(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match="cells mapping"):
        run_calc_suite(FIXTURE, {"oracle": "x", "cells": [1, 2]}, tmp_path)
    with pytest.raises(ValueError, match="has no value"):
        run_calc_suite(FIXTURE, _expected({"A1": {"type": "number"}}), tmp_path)


def test_calc_suite_engine_exception_is_contained(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        calc_module,
        "_engines",
        lambda: (_FakeEngine("boom", raises=RuntimeError("kaput")),),
    )
    results = run_calc_suite(FIXTURE, _expected({"A1": _cell(1)}), tmp_path)
    assert results["engines"]["boom"]["status"] == "failed"


# --------------------------------------------------------------------------
# Calc helper contracts
# --------------------------------------------------------------------------


def test_load_expected_accepts_dict_file_and_dir(tmp_path: Path) -> None:
    payload = {"oracle": "o", "cells": {"A1": _cell(1)}}
    assert calc_module._load_expected(payload) is payload

    file_path = tmp_path / "custom.json"
    file_path.write_text(json.dumps(payload))
    assert calc_module._load_expected(file_path) == payload

    (tmp_path / "expected_values.json").write_text(json.dumps(payload))
    assert calc_module._load_expected(tmp_path) == payload


def test_committed_calc_fixture_is_formula_dense_and_cache_free() -> None:
    references = formula_cells(FIXTURE)
    assert len(references) == 133

    cached = calc_module._read_values(FIXTURE, references[:10], reader="openpyxl")
    assert set(cached.values()) == {None}


def test_libreoffice_engine_reports_missing_binary(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(calc_module, "SOFFICE_PATH", Path("/nonexistent/soffice"))
    engine = LibreOfficeCalcEngine()
    assert engine.available() is False
    result = engine.calculate(FIXTURE, tmp_path / "out.xlsx")
    assert result["status"] == "unavailable"
    assert "not found" in result["reason"]

    reason = calc_module.recalculate_with_libreoffice(FIXTURE, tmp_path / "out.xlsx")
    assert reason is not None and "not found" in reason


def test_aspose_calc_engine_unavailable_without_package(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setitem(sys.modules, "aspose", None)
    assert AsposeCellsFossCalcEngine().available() is False


def _fake_oracle_catalog(available: bool) -> Any:
    class _Tool:
        def is_available(self) -> bool:
            return available

    class _Catalog:
        def __call__(self, repo_root: Path | None = None) -> dict[str, Any]:
            return {"zavora": _Tool()}

    return _Catalog()


def test_zavora_calc_engine_unavailable_without_helper(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        calc_module, "external_oracle_catalog", _fake_oracle_catalog(False)
    )
    engine = ZavoraCalcEngine()
    assert engine.available() is False
    result = engine.calculate(FIXTURE, tmp_path / "out.xlsx")
    assert result["status"] == "unavailable"
    assert result["reason"] == "Zavora helper is unavailable"


# --------------------------------------------------------------------------
# Mutation-suite control flow
# --------------------------------------------------------------------------


def _write_template_with_manifest(directory: Path) -> Path:
    template = directory / "template.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Inputs"
    sheet["A1"] = "original"
    sheet["B1"] = 1
    workbook.save(template)
    manifest = {
        "mutation": {"sheet": "Inputs", "cell": "A1", "value": "<engine>-landed"},
        "second_mutation": {"sheet": "Inputs", "cell": "B1", "value": 42},
    }
    (directory / "manifest.json").write_text(json.dumps(manifest))
    return template


class _UnavailableEngine:
    name = "ghost"

    def available(self) -> bool:
        return False

    def mutate(
        self, template: Path, output: Path, mutations: list[dict[str, Any]]
    ) -> None:
        raise AssertionError("unavailable engines must not run")


def test_run_mutation_suite_rejects_zero_repeats(tmp_path: Path) -> None:
    template = _write_template_with_manifest(tmp_path)
    with pytest.raises(ValueError, match="at least 1"):
        mutation_module.run_mutation_suite(template, tmp_path / "out", repeats=0)


def test_run_mutation_suite_unavailable_engine_row(tmp_path: Path) -> None:
    template = _write_template_with_manifest(tmp_path)
    monkey = pytest.MonkeyPatch()
    monkey.setattr(
        mutation_module, "modifiable_engines", lambda: [_UnavailableEngine()]
    )
    try:
        results = mutation_module.run_mutation_suite(
            template, tmp_path / "out", repeats=1
        )
    finally:
        monkey.undo()
    row = results["engines"]["ghost"]
    assert row["status"] == "unavailable"
    assert row["available"] is False
    assert row["wall_ms_median"] is None


def test_run_mutation_suite_reports_failed_engine(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    template = _write_template_with_manifest(tmp_path)

    def _explode(
        engine: Any,
        template: Path,
        output: Path,
        mutations: list[dict[str, Any]],
    ) -> tuple[float, float]:
        raise RuntimeError("driver crashed")

    monkeypatch.setattr(mutation_module, "_run_subprocess_mutation", _explode)
    monkeypatch.setattr(
        mutation_module,
        "modifiable_engines",
        lambda: [_UnavailableEngine.__new__(_UnavailableEngine)],  # noqa: E501
    )

    class _LiveEngine(_UnavailableEngine):
        name = "live"

        def available(self) -> bool:
            return True

    monkeypatch.setattr(mutation_module, "modifiable_engines", lambda: [_LiveEngine()])
    results = mutation_module.run_mutation_suite(template, tmp_path / "out", repeats=1)
    row = results["engines"]["live"]
    assert row["status"] == "failed"
    assert "driver crashed" in row["details"]["error"]
    assert row["preservation_score"] is None


def test_run_mutation_suite_passes_when_mutations_land(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    template = _write_template_with_manifest(tmp_path)

    def _apply(
        engine: Any,
        template: Path,
        output: Path,
        mutations: list[dict[str, Any]],
    ) -> tuple[float, float]:
        workbook = openpyxl.load_workbook(template)
        try:
            for mutation in mutations:
                sheet = workbook[str(mutation["sheet"])]
                sheet[str(mutation["cell"])] = mutation["value"]
            workbook.save(output)
        finally:
            workbook.close()
        return 12.5, 2048.0

    monkeypatch.setattr(mutation_module, "_run_subprocess_mutation", _apply)

    class _ApplyingEngine(_UnavailableEngine):
        name = "applying"

        def available(self) -> bool:
            return True

    monkeypatch.setattr(
        mutation_module, "modifiable_engines", lambda: [_ApplyingEngine()]
    )
    results = mutation_module.run_mutation_suite(template, tmp_path / "out", repeats=2)
    row = results["engines"]["applying"]
    assert row["status"] == "passed"
    assert row["wall_ms_median"] == 12.5
    assert row["peak_rss_kb_median"] == 2048.0
    assert row["preservation_score"] == 100.0
    assert row["details"]["mutations_landed"] is True


def test_run_mutation_suite_flags_integrity_when_values_do_not_land(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    template = _write_template_with_manifest(tmp_path)

    def _copy_only(
        engine: Any,
        template: Path,
        output: Path,
        mutations: list[dict[str, Any]],
    ) -> tuple[float, float]:
        import shutil

        shutil.copyfile(template, output)
        return 5.0, 1024.0

    monkeypatch.setattr(mutation_module, "_run_subprocess_mutation", _copy_only)

    class _LazyEngine(_UnavailableEngine):
        name = "lazy"

        def available(self) -> bool:
            return True

    monkeypatch.setattr(mutation_module, "modifiable_engines", lambda: [_LazyEngine()])
    results = mutation_module.run_mutation_suite(template, tmp_path / "out", repeats=1)
    row = results["engines"]["lazy"]
    assert row["status"] == "integrity-failed"
    assert row["preservation_score"] == 0.0
    assert "did not match" in row["details"]["integrity_failure"]


def test_mutations_for_engine_substitutes_placeholder() -> None:
    manifest = {
        "mutation": {"sheet": "S", "cell": "A1", "value": "<engine>-v"},
        "second_mutation": {"sheet": "S", "cell": "B1", "value": 7},
    }
    mutations = mutation_module._mutations_for_engine(manifest, "wolfxl")
    assert mutations[0]["value"] == "wolfxl-v"
    assert mutations[1] == {"sheet": "S", "cell": "B1", "value": 7}


def test_mutations_landed_rejects_unreadable_output(tmp_path: Path) -> None:
    bogus = tmp_path / "bogus.xlsx"
    bogus.write_bytes(b"not an xlsx")
    mutations = [{"sheet": "S", "cell": "A1", "value": 1}]
    assert mutation_module._mutations_landed(bogus, mutations) is False


def test_parse_time_output_extracts_elapsed_and_rss() -> None:
    stderr = "\n".join(
        [
            "        0.05 real",
            "        0.04 user",
            "        0.01 sys",
            "  4194304  maximum resident set size",
        ]
    )
    elapsed, rss = mutation_module._parse_time_output(stderr)
    assert elapsed == 0.05
    assert rss == 4096.0

    elapsed_none, rss_none = mutation_module._parse_time_output("garbage")
    assert elapsed_none is None and rss_none is None


def test_package_integrity_helpers_flag_damage() -> None:
    rels_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    parts = {
        "xl/_rels/workbook.xml.rels": (
            '<Relationships xmlns="'
            + rels_ns
            + '"><Relationship Id="rId1" '
            'Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            'relationships/worksheet" '
            'Target="worksheets/sheet1.xml"/></Relationships>'
        ).encode(),
        "xl/workbook.xml": b"<workbook/>",
        "[Content_Types].xml": (
            b'<Types xmlns="http://schemas.openxmlformats.org/package/2006/'
            b'content-types"><Default Extension="rels" ContentType="a"/></Types>'
        ),
    }
    dangling = mutation_module._dangling_relationships(parts)
    assert dangling == ["xl/_rels/workbook.xml.rels: xl/worksheets/sheet1.xml"]

    # workbook.xml has neither an Override nor an xml Default declaration.
    missing = mutation_module._missing_content_types(parts)
    assert "xl/workbook.xml" in missing

    absent_types = mutation_module._missing_content_types(
        {"xl/workbook.xml": b"<workbook/>"}
    )
    assert absent_types == ["[Content_Types].xml"]

    resolved = mutation_module._relationship_target(
        "xl/_rels/x.rels", "worksheets/s.xml"
    )
    assert resolved == "xl/worksheets/s.xml"
    worksheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    present = f'<x xmlns="{worksheet_ns}"><b/></x>'.encode()
    assert mutation_module._xml_has_element(present, "b") is True
    assert mutation_module._xml_has_element(b"<a><b/></a>", "b") is False
    assert mutation_module._xml_has_element(b"not xml <", "b") is False


# --------------------------------------------------------------------------
# Modifiable engine guards
# --------------------------------------------------------------------------


def test_mutation_target_splits_sheet_cell_value() -> None:
    sheet, cell, value = _mutation_target({"sheet": "Inputs", "cell": "A1", "value": 5})
    assert (sheet, cell, value) == ("Inputs", "A1", 5)


def test_openpyxl_engine_mutates_template(tmp_path: Path) -> None:
    template = _write_template_with_manifest(tmp_path)
    engine = OpenpyxlEngine()
    assert engine.available() is True
    output = tmp_path / "mutated.xlsx"
    engine.mutate(
        template,
        output,
        [
            {"sheet": "Inputs", "cell": "A1", "value": "changed"},
            {"sheet": "Inputs", "cell": "B1", "value": 9},
        ],
    )
    workbook = openpyxl.load_workbook(output)
    try:
        assert workbook["Inputs"]["A1"].value == "changed"
        assert workbook["Inputs"]["B1"].value == 9
    finally:
        workbook.close()


def test_optional_engines_report_absence(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(modifiable_module, "_aspose_cells_foss", None)
    engine = modifiable_module.AsposeFossEngine()
    assert engine.available() is False
    with pytest.raises(RuntimeError, match="not installed"):
        engine.mutate(Path("t.xlsx"), Path("o.xlsx"), [])

    monkeypatch.setattr(
        modifiable_module, "external_oracle_catalog", _fake_oracle_catalog(False)
    )
    assert modifiable_module.ZavoraEngine().available() is False


# --------------------------------------------------------------------------
# Renderers
# --------------------------------------------------------------------------


def test_mutation_report_renders_all_verdict_branches(tmp_path: Path) -> None:
    results: dict[str, Any] = {
        "metadata": {"template_sha256": "abc", "repeats": 1, "platform": "Darwin"},
        "engines": {
            "perfect": {
                "available": True,
                "status": "passed",
                "wall_ms_median": 100.0,
                "peak_rss_kb_median": 2048.0,
                "preservation_score": 100.0,
                "details": {},
            },
            "lossy": {
                "available": True,
                "status": "passed",
                "wall_ms_median": 200.0,
                "peak_rss_kb_median": 4096.0,
                "preservation_score": 60.0,
                "details": {},
            },
            "broken": {
                "available": True,
                "status": "integrity-failed",
                "wall_ms_median": 300.0,
                "peak_rss_kb_median": None,
                "preservation_score": 0.0,
                "details": {"integrity_failure": "nope"},
            },
            "crashed": {
                "available": True,
                "status": "failed",
                "wall_ms_median": None,
                "peak_rss_kb_median": None,
                "preservation_score": None,
                "details": {"error": "RuntimeError: x"},
            },
            "ghost": {
                "available": False,
                "status": "unavailable",
                "wall_ms_median": None,
                "peak_rss_kb_median": None,
                "preservation_score": None,
                "details": {},
            },
        },
    }
    render_mutation_report(results, tmp_path)
    persisted: dict[str, Any] = json.loads((tmp_path / "results.json").read_text())
    engines = persisted["engines"]
    assert isinstance(engines, dict)
    assert set(engines) == set(results["engines"])

    readme = (tmp_path / "README.md").read_text()
    for marker in ("Preserved", "Loss detected", "Mutation integrity failed", "Failed"):
        assert marker in readme, marker


# --------------------------------------------------------------------------
# Zavora adapter payload and failure contracts
# --------------------------------------------------------------------------


def test_zavora_adapter_reports_unavailable_helper(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(
        zavora_module, "external_oracle_catalog", _fake_oracle_catalog(False)
    )
    assert ZavoraAdapter.is_available() is False


def _fake_oracle_result(
    *, skipped: bool = False, passed: bool = False, payload: Any = None
) -> Any:
    class _Result:
        def __init__(self) -> None:
            self.skipped = skipped
            self.passed = passed
            self.payload = payload if isinstance(payload, dict) else {}
            self.notes = "helper skipped" if skipped else None
            self.stderr = ""
            self.stdout = ""

    return _Result()


def test_zavora_save_raises_when_helper_skips_or_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    adapter = ZavoraAdapter()
    monkeypatch.setattr(
        zavora_module,
        "run_external_oracle",
        lambda *a, **k: _fake_oracle_result(skipped=True),
    )
    with pytest.raises(FileNotFoundError, match="helper skipped"):
        adapter.save_workbook({"sheets": []}, tmp_path / "out.xlsx")

    monkeypatch.setattr(
        zavora_module,
        "run_external_oracle",
        lambda *a, **k: _fake_oracle_result(
            passed=False, payload={"message": "rust panic"}
        ),
    )
    with pytest.raises(RuntimeError, match="rust panic"):
        adapter.save_workbook({"sheets": []}, tmp_path / "out.xlsx")


def test_zavora_write_cell_value_payload_variants() -> None:
    adapter = ZavoraAdapter()
    workbook = adapter.create_workbook()
    adapter.write_cell_value(
        workbook, "S", "A1", CellValue(type=CellType.FORMULA, value=0, formula="=1+1")
    )
    adapter.write_cell_value(
        workbook, "S", "A2", CellValue(type=CellType.STRING, value="")
    )
    adapter.write_cell_value(
        workbook, "S", "A3", CellValue(type=CellType.NUMBER, value=7)
    )
    assert workbook["cells"][0] == {
        "sheet": "S",
        "cell": "A1",
        "type": "formula",
        "formula": "=1+1",
        "value": 0,
    }
    assert workbook["cells"][1]["type"] == "blank"
    assert workbook["cells"][1]["value"] is None
    assert workbook["cells"][2] == {
        "sheet": "S",
        "cell": "A3",
        "type": str(CellType.NUMBER),
        "value": 7,
    }
    assert {"name": "S"} in workbook["sheets"]


def test_zavora_conditional_format_mapping_helpers() -> None:
    adapter = ZavoraAdapter()
    assert adapter._map_conditional_format_type("colorScale") == "3_color_scale"
    assert adapter._map_conditional_format_type("dataBar") == "data_bar"
    assert adapter._map_conditional_format_type("expression") == "formula"
    assert adapter._map_conditional_format_type("cellIs") == "cell"
    assert (
        adapter._map_conditional_format_criteria("expression", None, "=Ref!$A$1>0")
        == "Ref!$A$1>0"
    )
    assert (
        adapter._map_conditional_format_criteria("cellIs", "greaterThan", None) == ">"
    )
    assert (
        adapter._map_conditional_format_criteria("cellIs", "unknownOperator", None)
        == ">"
    )


def test_zavora_read_operations_raise_structured_unsupported() -> None:
    adapter = ZavoraAdapter()
    workbook = adapter.create_workbook()
    with pytest.raises(NotImplementedError, match="write-only"):
        adapter.read_cell_value(workbook, "S", "A1")
    with pytest.raises(NotImplementedError, match="write-only"):
        adapter.get_sheet_names(workbook)
